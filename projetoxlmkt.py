import requests
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


# 1. Pegando os dados do site
link_api = "https://economia.awesomeapi.com.br/last/USD-BRL,EUR-BRL,BTC-BRL,XAU-BRL"
requisicao = requests.get(link_api)
dados = requisicao.json()

# 2. Criando o arquivo Excel
excel = Workbook()
aba = excel.active
aba.title = "Preços do Mercado"

# Estilos básicos para o cabeçalho
cor_fundo = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
fonte_branca = Font(bold=True, color="FFFFFF")
centralizar = Alignment(horizontal="center", vertical="center")

# 3. Criando os nomes das colunas
colunas = ["Nome", "Compra", "Venda", "Máximo", "Mínimo", "Variação %"]
aba.append(colunas)

for num, celula in enumerate(aba[1], 1):
    celula.fill = cor_fundo
    celula.font = fonte_branca
    celula.alignment = centralizar
    # Ajusta o tamanho da coluna
    aba.column_dimensions[get_column_letter(num)].width = 15

# 4. Preenchendo a tabela
linha = 2
for moeda in dados.values():
    nome = moeda['name'].split('/')[0]
    compra = float(moeda['bid'])
    venda = float(moeda['ask'])
    maximo = float(moeda['high'])
    minimo = float(moeda['low'])
    variacao = float(moeda['pctChange'])

    aba.append([nome, compra, venda, maximo, minimo, variacao])

    # Formata os números como dinheiro
    aba[f'B{linha}'].number_format = 'R$ #,##0.00'
    aba[f'C{linha}'].number_format = 'R$ #,##0.00'
    aba[f'D{linha}'].number_format = 'R$ #,##0.00'
    aba[f'E{linha}'].number_format = 'R$ #,##0.00'

    # Se subiu fica verde, se caiu fica vermelho
    if variacao > 0:
        aba[f'F{linha}'].font = Font(color="00B050", bold=True)
    else:
        aba[f'F{linha}'].font = Font(color="FF0000", bold=True)

    linha += 1

# 5. Calculando a média no final
linha_media = linha + 1
aba[f'E{linha_media}'] = "Média:"
aba[f'E{linha_media}'].font = Font(bold=True)
aba[f'F{linha_media}'] = f"=AVERAGE(F2:F{linha-1})"
aba[f'F{linha_media}'].font = Font(bold=True)

# 6. Salvando
nome_arquivo = "precos_mercado.xlsx"
excel.save(nome_arquivo)

print(f"o arquivo '{nome_arquivo}' foi salvo")
